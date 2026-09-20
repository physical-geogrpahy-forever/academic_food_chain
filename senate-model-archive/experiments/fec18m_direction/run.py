#!/usr/bin/env python3
import csv, io, math, re, urllib.request, urllib.parse, itertools
from pathlib import Path
from datetime import datetime, timezone
from difflib import SequenceMatcher
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'experiments/poll_direction/results/poll_fixed_predictions.csv'
DM=ROOT/'data/processed/core_v2r_headline_design_matrix_with_personal_vote_audit.csv'
OUT=ROOT/'experiments/fec18m_direction/results'
DOC=ROOT/'docs/history/updates/2026-09-21_0380_fec18m-fundraising-direction.md'
OUT.mkdir(parents=True,exist_ok=True); DOC.parent.mkdir(parents=True,exist_ok=True)

YEARS=[2014,2018,2022]
GAMMA=[0,0.5,1,2,3,4,5,6,8,10,12,15,20]
SIGNALS=['individual_share','receipts_share','cash_share','candidate_money_share']

def load(p):
    with p.open('r',encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def norm(s):
    s=(s or '').lower()
    s=s.replace(',',' ')
    s=re.sub(r'\b(jr|sr|ii|iii|iv)\b',' ',s)
    s=re.sub(r'[^a-z0-9 ]+',' ',s)
    return re.sub(r'\s+',' ',s).strip()
def num(v):
    if v is None:return 0.0
    if isinstance(v,(int,float)):return float(v)
    s=str(v).replace('$','').replace(',','').replace('(','-').replace(')','').strip()
    try:return float(s)
    except:return 0.0
def fetch(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 CoreV2R-FEC18M/1.0'})
    return urllib.request.urlopen(req,timeout=120).read()

def discover_xlsx(year):
    candidates=[
      f'https://www.fec.gov/resources/campaign-finance-statistics/{year}/tables/congressional/ConCand2_{year}_18m.xlsx',
      f'https://www.fec.gov/resources/campaign-finance-statistics/{year}/tables/congressional/ConCand2_{year}_18M.xlsx',
    ]
    for u in candidates:
        try:
            b=fetch(u)
            if b[:2]==b'PK':return u,b
        except:pass
    page=f'https://www.fec.gov/campaign-finance-data/congressional-candidate-data-summary-tables/?year={year}&segment=18'
    html=fetch(page).decode('utf-8','replace')
    links=re.findall(r'href=["\']([^"\']+\.xlsx)["\']',html,re.I)
    links=[urllib.parse.urljoin(page,x) for x in links]
    ranked=[u for u in links if f'ConCand2_{year}_' in u and ('18m' in u.lower() or '18M' in u)]
    ranked += [u for u in links if f'ConCand2_{year}_' in u and u not in ranked]
    for u in ranked:
        try:
            b=fetch(u)
            if b[:2]==b'PK':return u,b
        except:pass
    raise RuntimeError(f'Could not locate FEC Table 2 18M xlsx for {year}; links sample={links[:10]}')

def parse_workbook(year,b):
    wb=load_workbook(io.BytesIO(b),data_only=True,read_only=True)
    best=None
    for ws in wb.worksheets:
        vals=list(ws.iter_rows(values_only=True))
        for i,row in enumerate(vals[:40]):
            txt=[str(x).strip() if x is not None else '' for x in row]
            low=[x.lower() for x in txt]
            if any('candidate' in x for x in low) and any('receipt' in x for x in low):
                score=sum(bool(x) for x in txt)
                if best is None or score>best[0]:best=(score,ws.title,i,vals)
    if best is None:raise RuntimeError(f'No header row detected in {year} workbook')
    _,sheet,hi,vals=best
    rawheaders=[str(x).strip() if x is not None else '' for x in vals[hi]]
    headers=[];seen={}
    for j,h in enumerate(rawheaders):
        key=h or f'col{j}'
        if key in seen:seen[key]+=1;key=f'{key}_{seen[key]}'
        else:seen[key]=0
        headers.append(key)
    rows=[]
    for row in vals[hi+1:]:
        if not any(x is not None and str(x).strip() for x in row):continue
        d={headers[j]:(row[j] if j<len(row) else None) for j in range(len(headers))}
        rows.append(d)
    return sheet,headers,rows

def col(headers,patterns):
    for p in patterns:
        for h in headers:
            if p in h.lower():return h
    return None

fec_rows=[];source_meta=[]
for year in YEARS:
    url,b=discover_xlsx(year)
    sheet,headers,rr=parse_workbook(year,b)
    hc={
      'candidate':col(headers,['candidate name','candidate']),
      'state':col(headers,['state']),
      'party':col(headers,['party']),
      'receipts':col(headers,['total receipts','receipts']),
      'individual':col(headers,['individual']),
      'cash':col(headers,['cash on hand','cash']),
      'candidate_money':col(headers,['candidate contributions','contrib & loans','candidate contrib','loans from candidate'])
    }
    if not hc['candidate'] or not hc['state'] or not hc['receipts']:
        raise RuntimeError(f'{year} missing columns: {hc}; headers={headers}')
    for r in rr:
        name=str(r.get(hc['candidate']) or '').strip()
        state=str(r.get(hc['state']) or '').strip().upper()
        if not name or len(state)!=2:continue
        fec_rows.append({'year':year,'name':name,'state':state,'party':str(r.get(hc['party']) or '').strip(),
                         'receipts':num(r.get(hc['receipts'])),'individual':num(r.get(hc['individual'])) if hc['individual'] else 0.0,
                         'cash':num(r.get(hc['cash'])) if hc['cash'] else 0.0,
                         'candidate_money':num(r.get(hc['candidate_money'])) if hc['candidate_money'] else 0.0})
    source_meta.append({'year':year,'url':url,'sheet':sheet,'headers':' | '.join(headers)})

base=[r for r in load(BASE) if r['variant']=='w30_h14_all']
dm=load(DM);dby={r['race_id']:r for r in dm}
byys={}
for x in fec_rows:byys.setdefault((x['year'],x['state']),[]).append(x)

def match_candidate(year,state,name,party_side):
    q=norm(name);pool=byys.get((year,state),[])
    scored=[]
    for x in pool:
        n=norm(x['name'])
        s=1.0 if q==n else SequenceMatcher(None,q,n).ratio()
        # surname equality is a strong guard
        qlast=q.split()[-1] if q else '';nlast=n.split()[-1] if n else ''
        if qlast and nlast and qlast==nlast:s+=0.15
        p=x['party'].upper()
        if party_side=='D' and ('DEM' in p or p=='D'):s+=0.05
        if party_side=='R' and ('REP' in p or p=='R'):s+=0.05
        scored.append((s,x))
    scored.sort(key=lambda z:z[0],reverse=True)
    if not scored or scored[0][0]<0.80:return None,0.0
    if len(scored)>1 and scored[0][0]-scored[1][0]<0.03:return None,scored[0][0]
    return scored[0][1],scored[0][0]

rows=[];review=[]
for r in base:
    cyc=int(r['test_cycle']);d=dby[r['race_id']]
    dn=d['d_side_candidate'];rn=d['r_side_candidate'];st=r['state_abbrev']
    dmch,ds=match_candidate(cyc,st,dn,'D');rmch,rs=match_candidate(cyc,st,rn,'R')
    review.append({'cycle':cyc,'state':st,'race_id':r['race_id'],'d_candidate':dn,'d_match':'' if not dmch else dmch['name'],'d_score':ds,
                   'r_candidate':rn,'r_match':'' if not rmch else rmch['name'],'r_score':rs})
    if not dmch or not rmch:continue
    def share(key):
        a=max(dmch[key],0.0);b=max(rmch[key],0.0)
        return 0.0 if a+b<=0 else (a-b)/(a+b)
    rows.append({'cycle':cyc,'race_id':r['race_id'],'state':st,'actual':float(r['actual']),'posterior':float(r['posterior']),
                 'individual_share':share('individual'),'receipts_share':share('receipts'),'cash_share':share('cash'),
                 'candidate_money_share':share('candidate_money')})

def evalset(dat,signal,gamma):
    ok=flips=0
    for r in dat:
        q=r['posterior']+gamma*r[signal]
        ok+=int((q>0)==(r['actual']>0));flips+=int((q>0)!=(r['posterior']>0))
    return ok,flips

choices=[];pred=[]
for tc in YEARS:
    train=[r for r in rows if r['cycle']<tc];test=[r for r in rows if r['cycle']==tc]
    if not train:
        signal='individual_share';gamma=0;ok=0;flips=0
    else:
        cand=[]
        for signal,gamma in itertools.product(SIGNALS,GAMMA):
            ok,flips=evalset(train,signal,gamma)
            cand.append((-ok,flips,gamma,signal))
        cand.sort(key=lambda z:(z[0],z[1],z[2],z[3]))
        z=cand[0];ok=-z[0];flips=z[1];gamma=z[2];signal=z[3]
    choices.append({'test_cycle':tc,'signal':signal,'gamma':gamma,'train_n':len(train),
                    'train_correct':ok if train else '','train_accuracy_pct':100*ok/len(train) if train else '',
                    'train_flips':flips,'test_finance_covered':len(test)})
    for r in test:
        q=r['posterior']+gamma*r[signal]
        pred.append({**r,'selected_signal':signal,'gamma':gamma,'adjusted_posterior':q,
                     'baseline_correct':int((r['posterior']>0)==(r['actual']>0)),
                     'adjusted_correct':int((q>0)==(r['actual']>0))})

# For unmatched races, preserve baseline rather than dropping them.
pmap={(r['cycle'],r['race_id']):r for r in pred}
allpred=[]
for b in base:
    key=(int(b['test_cycle']),b['race_id'])
    if key in pmap:allpred.append(pmap[key]);continue
    act=float(b['actual']);post=float(b['posterior'])
    allpred.append({'cycle':key[0],'race_id':key[1],'state':b['state_abbrev'],'actual':act,'posterior':post,
                    'individual_share':'','receipts_share':'','cash_share':'','candidate_money_share':'',
                    'selected_signal':'UNMATCHED_BASELINE','gamma':0,'adjusted_posterior':post,
                    'baseline_correct':int((post>0)==(act>0)),'adjusted_correct':int((post>0)==(act>0))})

summary=[]
for scope in ['combined','2014','2018','2022']:
    rr=allpred if scope=='combined' else [r for r in allpred if r['cycle']==int(scope)]
    n=len(rr);b=sum(r['baseline_correct'] for r in rr);a=sum(r['adjusted_correct'] for r in rr)
    summary.append({'scope':scope,'n':n,'finance_covered':sum(r['selected_signal']!='UNMATCHED_BASELINE' for r in rr),
                    'baseline_correct':b,'baseline_accuracy_pct':100*b/n,'adjusted_correct':a,
                    'adjusted_accuracy_pct':100*a/n,'net_correct_gain':a-b})

for fn,data in [('fec18m_summary.csv',summary),('fec18m_choices.csv',choices),('fec18m_predictions.csv',allpred),('fec18m_match_review.csv',review),('fec18m_source_meta.csv',source_meta)]:
    with (OUT/fn).open('w',encoding='utf-8',newline='') as f:
        fields=list(data[0].keys());w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(data)

changed=[r for r in allpred if r['baseline_correct']!=r['adjusted_correct']]
lines=['# FEC 18-month fundraising direction experiment','',
'- Generated UTC: '+datetime.now(timezone.utc).isoformat(),
'- FEC Table 2 uses campaign finance activity through June 30 of each election year.',
'- The same 18-month filing stage is used for 2014, 2018, 2022, before the 45-day polling snapshot.',
'- Tested signals: individual-contribution share, total-receipts share, cash-on-hand share, candidate contributions/loans share.',
'- For 2018, signal and gamma are selected using 2014 only; for 2022, using 2014+2018 only.',
'- Races without a confident FEC candidate match keep the 94/99 baseline prediction unchanged.','',
'## Result','',
'| scope | N | finance-covered | baseline correct | baseline acc | adjusted correct | adjusted acc | net |',
'|---|---:|---:|---:|---:|---:|---:|---:|']
for r in summary:lines.append(f"| {r['scope']} | {r['n']} | {r['finance_covered']} | {r['baseline_correct']} | {r['baseline_accuracy_pct']:.1f}% | {r['adjusted_correct']} | {r['adjusted_accuracy_pct']:.1f}% | {r['net_correct_gain']:+d} |")
lines += ['','## Selected finance signal','',
'| test | signal | gamma | train N | train accuracy | test finance coverage |',
'|---:|---|---:|---:|---:|---:|']
for r in choices:
    ta='NA' if r['train_accuracy_pct']=='' else f"{r['train_accuracy_pct']:.1f}%"
    lines.append(f"| {r['test_cycle']} | {r['signal']} | {r['gamma']} | {r['train_n']} | {ta} | {r['test_finance_covered']} |")
lines += ['','## Correctness changes','']
for r in changed:lines.append(f"- {r['cycle']} {r['state']} {r['race_id']}: baseline {'correct' if r['baseline_correct'] else 'wrong'} -> finance {'correct' if r['adjusted_correct'] else 'wrong'}")
lines += ['','## Decision','',
'Adopt only if combined direction accuracy exceeds 94/99 and candidate matching coverage is adequate.','',
'## Outputs','',
'- experiments/fec18m_direction/results/fec18m_summary.csv',
'- experiments/fec18m_direction/results/fec18m_choices.csv',
'- experiments/fec18m_direction/results/fec18m_predictions.csv',
'- experiments/fec18m_direction/results/fec18m_match_review.csv',
'- experiments/fec18m_direction/results/fec18m_source_meta.csv']
DOC.write_text('\n'.join(lines)+'\n',encoding='utf-8');print('\n'.join(lines))
