#!/usr/bin/env python3
import csv, io, re, math, urllib.request
from collections import defaultdict
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SEN=ROOT/'data/raw/fivethirtyeight/election_results_senate_2026-09-21.csv'
GOV=ROOT/'data/processed/source_snapshots/election_results_gubernatorial_d7a7cff101da.csv'
LEAN=ROOT/'data/processed/presidential_state_lean.csv'
STRUCT=ROOT/'data/snapshots/2026_structural_input_matrix_45d.csv'
EXP=ROOT/'data/snapshots/2026_candidate_experience_core.csv'
POLL=ROOT/'data/snapshots/2026_poll_45d_recency_proxy.csv'
OUTDIR=ROOT/'data/snapshots'
DOC=ROOT/'docs/analysis/2026_PRODUCTION_INPUT_MATRIX_45D.md'
OUTDIR.mkdir(parents=True,exist_ok=True);DOC.parent.mkdir(parents=True,exist_ok=True)

CANDIDATES={
 'AK':('Mary Peltola','Dan Sullivan'),
 'GA':('Jon Ossoff','Mike Collins'),
 'IA':('Josh Turek','Ashley Hinson'),
 'ME':('Troy Jackson','Susan Collins'),
 'MI':('Abdul El-Sayed','Mike Rogers'),
 'NH':('Chris Pappas','John Sununu'),
 'NC':('Roy Cooper','Michael Whatley'),
 'OH':('Sherrod Brown','Jon Husted'),
 'TX':('James Talarico','Ken Paxton'),
}

def load(p):
    with p.open('r',encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def norm(s):
    s=(s or '').strip().lower()
    if ',' in s:
        a,b=s.split(',',1);s=b+' '+a
    s=re.sub(r'\b(jr|sr|ii|iii|iv|mr|mrs|ms|dr|sen|rep|gov)\b',' ',s)
    s=re.sub(r'[^a-z0-9 ]+',' ',s)
    return re.sub(r'\s+',' ',s).strip()
def short_match(a,b):
    a=norm(a);b=norm(b)
    if a==b:return 1.0
    at=a.split();bt=b.split()
    if not at or not bt:return 0.0
    surname=1.0 if at[-1]==bt[-1] else 0.0
    first=1.0 if at[0]==bt[0] or at[0][0]==bt[0][0] else 0.0
    seq=SequenceMatcher(None,a,b).ratio()
    return .55*seq+.30*surname+.15*first
def candidate_entities(rows):
    out=[]
    seen=set()
    for r in rows:
        name=(r.get('candidate_name') or '').strip()
        pid=(r.get('politician_id') or '').strip()
        if not name or not pid:continue
        key=(pid,norm(name))
        if key in seen:continue
        seen.add(key)
        out.append({'pid':pid,'name':name,'state':r.get('state_abbrev','')})
    return out

sen=load(SEN);gov=load(GOV);leanrows=load(LEAN)
entities=candidate_entities(sen+gov)

# Candidate -> historical politician_id mapping, audited by score.
candidate_pid={}
pid_audit=[]
for st,(dn,rn) in CANDIDATES.items():
    for side,name in [('D',dn),('R',rn)]:
        pool=[x for x in entities if x['state']==st]
        scored=sorted([(short_match(name,x['name']),x) for x in pool],key=lambda z:z[0],reverse=True)
        best=scored[0] if scored else (0,None)
        pid=best[1]['pid'] if best[0]>=.76 else ''
        candidate_pid[(st,side)]=pid
        pid_audit.append({
          'state_abbrev':st,'side':side,'target_name':name,'matched_name':'' if not best[1] else best[1]['name'],
          'politician_id':pid,'match_score':best[0],
          'second_score':scored[1][0] if len(scored)>1 else ''
        })

# Historical PVI function.
lean={}
years=set()
for r in leanrows:
    try:y=int(r['cycle']);v=float(r['lean_vs_national_pctpt'])
    except:continue
    lean[(y,r['state_abbrev'])]=v;years.add(y)
years=sorted(years)
def pvi_state(cyc,st):
    ys=[y for y in years if y<cyc and (y,st) in lean]
    if len(ys)<2:return None
    return .67*lean[(ys[-1],st)]+.33*lean[(ys[-2],st)]

# Parse one D-vs-R race from election-results rows.
def race_margin(rr):
    # highest RCV round only
    nums=[]
    for x in rr:
        v=(x.get('ranked_choice_round') or '').strip()
        if v:
            try:nums.append(int(float(v)))
            except:pass
    if nums:
        mx=max(nums)
        filt=[]
        for x in rr:
            v=(x.get('ranked_choice_round') or '').strip()
            try:rv=int(float(v)) if v else None
            except:rv=None
            if rv==mx:filt.append(x)
        rr=filt
    cand={}
    for r in rr:
        key=(r.get('politician_id') or r.get('candidate_id') or r.get('candidate_name') or '').strip()
        if not key:continue
        c=cand.setdefault(key,{'pid':(r.get('politician_id') or '').strip(),'name':r.get('candidate_name') or '',
                               'votes':0,'parties':set(),'missing':False})
        for p in ((r.get('ballot_party') or ''),(r.get('party') or '')):
            if p.strip():c['parties'].add(p.strip().upper())
        v=(r.get('votes') or '').strip()
        if not v:c['missing']=True
        else:
            try:c['votes']+=int(float(v))
            except:c['missing']=True
    ds=[c for c in cand.values() if 'DEM' in c['parties']]
    rs=[c for c in cand.values() if 'REP' in c['parties']]
    if len(ds)!=1 or len(rs)!=1 or ds[0]['missing'] or rs[0]['missing'] or ds[0]['votes']+rs[0]['votes']<=0:return None
    return 100*(ds[0]['votes']-rs[0]['votes'])/(ds[0]['votes']+rs[0]['votes']),ds[0],rs[0]

# Build Senate and governor race residuals = D-R margin - PVI - cycle mean residual.
def build_perf(rawrows,office_kind):
    grouped=defaultdict(list)
    for r in rawrows:
        if (r.get('stage') or '').lower()!='general':continue
        grouped[(int(r['cycle']),str(r['race_id']))].append(r)
    races=[]
    for (cyc,rid),rr in grouped.items():
        q=race_margin(rr)
        if q is None:continue
        m,d,r=q;st=rr[0]['state_abbrev'];pv=pvi_state(cyc,st)
        if pv is None:continue
        races.append({'cycle':cyc,'race_id':rid,'state':st,'margin':m,'pvi':pv,'raw_resid':m-pv,
                      'd_pid':d['pid'],'r_pid':r['pid'],'d_name':d['name'],'r_name':r['name'],'office':office_kind})
    means={}
    for cyc in sorted(set(x['cycle'] for x in races)):
        vals=[x['raw_resid'] for x in races if x['cycle']==cyc]
        if vals:means[cyc]=sum(vals)/len(vals)
    bypid=defaultdict(list)
    for x in races:
        adj=x['raw_resid']-means[x['cycle']]
        if x['d_pid']:bypid[x['d_pid']].append({**x,'score':adj,'side':'D'})
        if x['r_pid']:bypid[x['r_pid']].append({**x,'score':-adj,'side':'R'})
    return bypid

senperf=build_perf(sen,'Senate')
govperf=build_perf(gov,'Governor')

def personal_for(pid,target_cycle=2026):
    vals=[]
    for store in (senperf,govperf):
        for x in store.get(pid,[]):
            if x['cycle']<target_cycle:
                vals.append(x)
    vals=sorted(vals,key=lambda x:(x['cycle'],x['office'],x['race_id']))
    if not vals:return {'raw_mean':0.0,'count':0,'shrunk_k2':0.0,'detail':''}
    mean=sum(x['score'] for x in vals)/len(vals)
    n=len(vals)
    shr=mean*n/(n+2.0)
    detail=' | '.join(f"{x['cycle']} {x['office']} score={x['score']:.2f}" for x in vals)
    return {'raw_mean':mean,'count':n,'shrunk_k2':shr,'detail':detail}

personal=[]
for st,(dn,rn) in CANDIDATES.items():
    dp=personal_for(candidate_pid[(st,'D')]) if candidate_pid[(st,'D')] else {'raw_mean':0,'count':0,'shrunk_k2':0,'detail':''}
    rp=personal_for(candidate_pid[(st,'R')]) if candidate_pid[(st,'R')] else {'raw_mean':0,'count':0,'shrunk_k2':0,'detail':''}
    personal.append({
      'state_abbrev':st,'candidate_D':dn,'candidate_R':rn,
      'd_politician_id':candidate_pid[(st,'D')],'r_politician_id':candidate_pid[(st,'R')],
      'd_personal_mean_overperf_pctpt':dp['raw_mean'],'d_personal_prior_statewide_count':dp['count'],'d_personal_shrunk_k2_pctpt':dp['shrunk_k2'],
      'r_personal_mean_overperf_pctpt':rp['raw_mean'],'r_personal_prior_statewide_count':rp['count'],'r_personal_shrunk_k2_pctpt':rp['shrunk_k2'],
      'PersonalVoteDiff_shrunk_D_minus_R_pctpt':dp['shrunk_k2']-rp['shrunk_k2'],
      'd_history':dp['detail'],'r_history':rp['detail']
    })

# FEC June-30 18m Top-50 2026.
def get(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 CoreV2R-2026-FEC/1.0'})
    return urllib.request.urlopen(req,timeout=180).read()
def num(v):
    if isinstance(v,(int,float)):return float(v)
    s=str(v or '').replace('$','').replace(',','').replace('(','-').replace(')','').strip()
    try:return float(s)
    except:return None
def parse_top50(metric,table):
    urls=[
      f'https://www.fec.gov/resources/campaign-finance-statistics/2026/tables/congressional/ConCand{table}_2026_18m.xlsx',
      f'https://www.fec.gov/resources/campaign-finance-statistics/2026/tables/congressional/ConCand{table}_2026_18M.xlsx'
    ]
    raw=None;used=''
    for u in urls:
        try:
            b=get(u)
            if b[:2]==b'PK':raw=b;used=u;break
        except Exception:pass
    if raw is None:return [],''
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    found=[]
    for ws in wb.worksheets:
        vals=list(ws.iter_rows(values_only=True));hi=None
        for i,row in enumerate(vals[:35]):
            low=[str(x or '').strip().lower() for x in row]
            if any('candidate name' in x for x in low) and any(x=='state' for x in low):
                hi=i;break
        if hi is None:continue
        headers=[str(x or '').strip() for x in vals[hi]]
        def idx(pats):
            for p in pats:
                for j,h in enumerate(headers):
                    if p in h.lower():return j
            return None
        ni=idx(['candidate name','candidate']);si=idx(['state']);pi=idx(['party'])
        ai=None
        for j in range(len(headers)-1,-1,-1):
            h=headers[j].lower()
            if any(k in h for k in ['receipt','individual','cash on hand','contribution']):ai=j;break
        if None in (ni,si,ai):continue
        for row in vals[hi+1:]:
            name=str(row[ni] or '').strip() if ni<len(row) else '';st=str(row[si] or '').strip().upper() if si<len(row) else ''
            amount=num(row[ai] if ai<len(row) else None)
            if not name or len(st)!=2 or amount is None:continue
            found.append({'metric':metric,'name':name,'state':st,'party':str(row[pi] or '').strip() if pi is not None and pi<len(row) else '',
                          'amount':amount,'source_url':used})
    return found,used

tables={'receipts':'6a','individual':'6b','cash':'6f'}
fecrows=[];fecsrc=[]
for metric,t in tables.items():
    rr,u=parse_top50(metric,t);fecrows+=rr;fecsrc.append({'metric':metric,'url':u,'rows':len(rr)})
fecindex=defaultdict(list)
for r in fecrows:fecindex[(r['metric'],r['state'])].append(r)

def match_fec(metric,st,name,side):
    q=norm(name);qt=q.split();sc=[]
    for x in fecindex.get((metric,st),[]):
        n=norm(x['name']);nt=n.split()
        seq=1.0 if q==n else SequenceMatcher(None,q,n).ratio()
        surname=1 if qt and nt and qt[-1]==nt[-1] else 0
        first=1 if qt and nt and (qt[0]==nt[0] or qt[0][0]==nt[0][0]) else 0
        party=(x['party'] or '').upper()
        pb=.06 if (side=='D' and ('DEM' in party or 'DFL' in party)) or (side=='R' and 'REP' in party) else 0
        score=.55*seq+.25*surname+.14*first+pb
        sc.append((score,x))
    sc.sort(key=lambda z:z[0],reverse=True)
    if not sc or sc[0][0]<.70:return None,(sc[0][0] if sc else 0)
    if len(sc)>1 and sc[0][0]-sc[1][0]<.03:return None,sc[0][0]
    return sc[0][1],sc[0][0]

fec=[]
fec_review=[]
for st,(dn,rn) in CANDIDATES.items():
    rec={'state_abbrev':st,'candidate_D':dn,'candidate_R':rn}
    for metric in tables:
        dm,ds=match_fec(metric,st,dn,'D');rm,rs=match_fec(metric,st,rn,'R')
        fec_review.append({'state_abbrev':st,'metric':metric,'d_candidate':dn,'d_match':'' if dm is None else dm['name'],'d_score':ds,
                           'r_candidate':rn,'r_match':'' if rm is None else rm['name'],'r_score':rs})
        rec['d_'+metric]='' if dm is None else dm['amount']
        rec['r_'+metric]='' if rm is None else rm['amount']
        if dm is not None and rm is not None and dm['amount']+rm['amount']>0:
            rec[metric+'_share_D_minus_R']=(dm['amount']-rm['amount'])/(dm['amount']+rm['amount'])
        else:rec[metric+'_share_D_minus_R']=''
    fec.append(rec)

# Merge final production input matrix.
struct={r['state_abbrev']:r for r in load(STRUCT)}
exp={r['state_abbrev']:r for r in load(EXP)}
poll={r['state_abbrev']:r for r in load(POLL)}
pers={r['state_abbrev']:r for r in personal}
fin={r['state_abbrev']:r for r in fec}

final=[]
for st in CANDIDATES:
    s=struct[st];e=exp[st];p=poll[st];v=pers[st];f=fin[st]
    final.append({
      'snapshot_date':'2026-09-19','state_abbrev':st,'candidate_D':s['candidate_D'],'candidate_R':s['candidate_R'],
      'pvi_D_minus_R':s['pvi_2026_d_minus_r_pctpt'],'same_seat_D_minus_R':s['same_seat_d_minus_r_pctpt'],
      'incumbency_diff_D_minus_R':s['incumbency_diff_D_minus_R'],'outparty_incumbent_signed':s['outparty_incumbent_signed'],
      'SenateExperienceDiff_D_minus_R':e['SenateExperienceDiff_D_minus_R'],
      'GovernorExperienceDiff_D_minus_R':e['GovernorExperienceDiff_D_minus_R'],
      'HouseExperienceDiff_D_minus_R':e['HouseExperienceDiff_D_minus_R'],
      'PersonalVoteDiff_shrunk_D_minus_R':v['PersonalVoteDiff_shrunk_D_minus_R_pctpt'],
      'relative_economic_growth_presparty_signed':s['relative_economic_growth_presparty_signed'],
      'national_generic_ballot_D_minus_R':s['national_generic_ballot_D_minus_R_pctpt'],
      'poll_D_minus_R_recency_proxy':p['poll_D_minus_R_recency_proxy'],
      'production_override_trigger':p['production_override_trigger'],
      'fec_receipts_share_D_minus_R':f['receipts_share_D_minus_R'],
      'fec_individual_share_D_minus_R':f['individual_share_D_minus_R'],
      'fec_cash_share_D_minus_R':f['cash_share_D_minus_R'],
      'third_party_uncertainty_flag':1 if st=='AK' else 0,
      'exact_cutoff':'2026-09-19 US date'
    })

for path,arr in [
 (OUTDIR/'2026_personal_vote_core.csv',personal),
 (OUTDIR/'2026_fec_june30_top50.csv',fec),
 (OUTDIR/'2026_fec_june30_match_review.csv',fec_review),
 (OUTDIR/'2026_fec_june30_source_meta.csv',fecsrc),
 (OUTDIR/'2026_candidate_pid_match_audit.csv',pid_audit),
 (OUTDIR/'2026_production_input_matrix_45d.csv',final)
]:
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(arr[0].keys()));w.writeheader();w.writerows(arr)

lines=['# 2026 production input matrix at exact 45-day cutoff','',
'- Cutoff: 2026-09-19 US date.',
'- PersonalVote uses prior Senate and Governor general elections only.',
'- Candidate-side overperformance = D/R-side residual from state PVI after subtracting the same-cycle mean residual.',
'- PersonalVote is shrunk with k=2 for the audit column.',
'- FEC finance uses official 2026 18-month Top-50 tables through June 30. Missing Top-50 membership is not converted to zero.',
'- FEC remains an auxiliary diagnostic because it was not adopted as a production direction modifier in historical OOS.',
'- Poll recency values are cutoff diagnostics with equal-sample assumption because RCP feed does not provide sample sizes.','',
'## Candidate PersonalVote','',
'| State | D PV | n | R PV | n | D-R shrunk PV |',
'|---|---:|---:|---:|---:|---:|']
for r in personal:
    lines.append(f"| {r['state_abbrev']} | {r['d_personal_mean_overperf_pctpt']:.2f} | {r['d_personal_prior_statewide_count']} | {r['r_personal_mean_overperf_pctpt']:.2f} | {r['r_personal_prior_statewide_count']} | {r['PersonalVoteDiff_shrunk_D_minus_R_pctpt']:.2f} |")
lines += ['','## Complete input rows','',
'The full machine-readable matrix is data/snapshots/2026_production_input_matrix_45d.csv.','',
'## Important separation','',
'- Locked model inputs: PVI, SameSeat, incumbency, OutPartyIncumbent, experience, PersonalVote, RelativeEconomicGrowth, National, 45-day poll layer.',
'- Auxiliary diagnostics only: FEC finance shares and third-party uncertainty flag.',
'- No 2026 election outcome is used anywhere in these inputs.']
DOC.write_text('\n'.join(lines)+'\n',encoding='utf-8')
print('\n'.join(lines))
