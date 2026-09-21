#!/usr/bin/env python3
import csv,io,re,urllib.request
from collections import defaultdict
from pathlib import Path
from difflib import SequenceMatcher
from datetime import datetime,timezone
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data/snapshots/2026_fec_june30_top50.csv'
META=ROOT/'data/snapshots/2026_fec_june30_source_meta.csv'
DOC=ROOT/'docs/analysis/2026_FEC_JUNE30_BLOCK.md'

TARGET=[
('AK','D','Mary Peltola'),('AK','R','Dan Sullivan'),('GA','D','Jon Ossoff'),('GA','R','Mike Collins'),
('IA','D','Josh Turek'),('IA','R','Ashley Hinson'),('ME','D','Troy Jackson'),('ME','R','Susan Collins'),
('MI','D','Abdul El-Sayed'),('MI','R','Mike Rogers'),('NH','D','Chris Pappas'),('NH','R','John Sununu'),
('NC','D','Roy Cooper'),('NC','R','Michael Whatley'),('OH','D','Sherrod Brown'),('OH','R','Jon Husted'),
('TX','D','James Talarico'),('TX','R','Ken Paxton')]
TABLES={'receipts':'6a','individual':'6b','cash':'6f'}

def norm(s):
    s=(s or '').strip().lower()
    if ',' in s:
        a,b=s.split(',',1);s=b+' '+a
    s=re.sub(r'\b(jr|sr|ii|iii|iv|mr|mrs|ms|dr|sen)\b',' ',s)
    s=re.sub(r'[^a-z0-9 ]+',' ',s)
    return re.sub(r'\s+',' ',s).strip()
def num(v):
    if isinstance(v,(int,float)):return float(v)
    s=str(v or '').replace('$','').replace(',','').replace('(','-').replace(')','').strip()
    try:return float(s)
    except:return None
def get(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 CoreV2R-2026-FEC/1.0'})
    return urllib.request.urlopen(req,timeout=120).read()
def parse(metric,table):
    urls=[f'https://www.fec.gov/resources/campaign-finance-statistics/2026/tables/congressional/ConCand{table}_2026_18m.xlsx',
          f'https://www.fec.gov/resources/campaign-finance-statistics/2026/tables/congressional/ConCand{table}_2026_18M.xlsx']
    raw=None;used=None
    for u in urls:
        try:
            b=get(u)
            if b[:2]==b'PK':raw=b;used=u;break
        except Exception:pass
    if raw is None:return [],'',False
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True);found=[]
    for ws in wb.worksheets:
        vals=list(ws.iter_rows(values_only=True));hi=None
        for i,row in enumerate(vals[:30]):
            low=[str(x or '').strip().lower() for x in row]
            if any('candidate name' in x for x in low) and any(x=='state' for x in low):hi=i;break
        if hi is None:continue
        headers=[str(x or '').strip() for x in vals[hi]]
        def idx(parts):
            for p in parts:
                for j,h in enumerate(headers):
                    if p in h.lower():return j
            return None
        ni=idx(['candidate name','candidate']);si=idx(['state']);pi=idx(['party']);ai=None
        for j in range(len(headers)-1,-1,-1):
            h=headers[j].lower()
            if any(k in h for k in ['receipt','individual','cash on hand','contribution']):ai=j;break
        if None in (ni,si,ai):continue
        for row in vals[hi+1:]:
            name=str(row[ni] or '').strip() if ni<len(row) else ''
            st=str(row[si] or '').strip().upper() if si<len(row) else ''
            if not name or len(st)!=2:continue
            amount=num(row[ai] if ai<len(row) else None)
            if amount is None:continue
            party=str(row[pi] or '').strip() if pi is not None and pi<len(row) else ''
            found.append({'metric':metric,'name':name,'state':st,'party':party,'amount':amount,'source_url':used})
    return found,used,True

allrows=[];meta=[]
for metric,table in TABLES.items():
    rows,url,ok=parse(metric,table);allrows+=rows
    meta.append({'metric':metric,'available_18m':ok,'source_url':url,'parsed_rows':len(rows)})
index=defaultdict(list)
for r in allrows:index[(r['metric'],r['state'])].append(r)

def party_bonus(p,side):
    p=(p or '').upper()
    if side=='D':return .06 if ('DEM' in p or 'DFL' in p) else 0
    return .06 if 'REP' in p else 0
def match(metric,st,name,side):
    q=norm(name);qt=q.split();sc=[]
    for x in index.get((metric,st),[]):
        n=norm(x['name']);nt=n.split()
        seq=1 if q==n else SequenceMatcher(None,q,n).ratio()
        sur=1 if qt and nt and qt[-1]==nt[-1] else 0
        fst=1 if qt and nt and (qt[0]==nt[0] or qt[0][0]==nt[0][0]) else 0
        jac=len(set(qt)&set(nt))/max(len(set(qt)|set(nt)),1)
        score=.5*seq+.25*sur+.15*fst+.1*jac+party_bonus(x['party'],side)
        sc.append((score,x))
    sc.sort(key=lambda z:z[0],reverse=True)
    if not sc or sc[0][0]<.70:return None,(sc[0][0] if sc else 0)
    if len(sc)>1 and sc[0][0]-sc[1][0]<.035:return None,sc[0][0]
    return sc[0][1],sc[0][0]

out=[]
for st in sorted(set(x[0] for x in TARGET)):
    rec={'state_abbrev':st}
    for side in ['D','R']:
        name=next(n for s,si,n in TARGET if s==st and si==side);rec['candidate_'+side]=name
        for metric in TABLES:
            m,score=match(metric,st,name,side)
            rec[f'{side.lower()}_{metric}']='' if not m else m['amount']
            rec[f'{side.lower()}_{metric}_match']='' if not m else m['name']
            rec[f'{side.lower()}_{metric}_score']=score
    for metric in TABLES:
        a=rec[f'd_{metric}'];b=rec[f'r_{metric}']
        rec[metric+'_share_D_minus_R']='' if a=='' or b=='' or float(a)+float(b)==0 else (float(a)-float(b))/(float(a)+float(b))
    out.append(rec)

for path,arr in [(OUT,out),(META,meta)]:
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(arr[0].keys()));w.writeheader();w.writerows(arr)

lines=['# 2026 FEC June-30 18-month block','',
'- Generated UTC: '+datetime.now(timezone.utc).isoformat(),
'- Required cutoff: activity through 2026-06-30.',
'- Source family: official FEC congressional candidate statistical tables 6a, 6b, 6f.',
'- Missing Top-50 membership is retained as missing, never zero.','',
'## Source availability','']
for m in meta:lines.append(f"- {m['metric']}: available={m['available_18m']}, parsed_rows={m['parsed_rows']}, source={m['source_url'] or 'NA'}")
lines += ['','## Race-level coverage','',
'| State | receipts share D-R | individual share D-R | cash share D-R |','|---|---:|---:|---:|']
for r in out:
    def fmt(v):return 'NA' if v=='' else f"{float(v):.3f}"
    lines.append(f"| {r['state_abbrev']} | {fmt(r['receipts_share_D_minus_R'])} | {fmt(r['individual_share_D_minus_R'])} | {fmt(r['cash_share_D_minus_R'])} |")
lines += ['','## Interpretation','',
'- Finance is an auxiliary input only. The locked historical production model did not promote FEC finance to a universal structural term.',
'- These values are retained for race diagnostics and for the narrow finance-aware post-hoc candidate analysis already documented.',
'- Do not infer a zero financial position from a missing Top-50 match.']
DOC.write_text('\n'.join(lines)+'\n',encoding='utf-8')
print('\n'.join(lines))
